#!/usr/bin/env python3
"""Semantically diff two .pptx files to detect a user's manual edits before rebuilding.
PowerPoint reorders XML on save, so compare MEANING (text runs, geometry, fills, fonts),
not raw XML. Usage: python3 diff_pptx.py <build.pptx> <delivered.pptx>"""
import sys
from pptx import Presentation

def load(path):
    p = Presentation(path); out = []
    for sl in p.slides:
        shapes = []
        for s in sl.shapes:
            try: box = (round(s.left/9525), round(s.top/9525), round(s.width/9525), round(s.height/9525))
            except Exception: box = (None,)*4
            fill = ""
            try:
                if s.fill.type == 1: fill = str(s.fill.fore_color.rgb)
            except Exception: pass
            fonts = []
            if s.has_text_frame:
                for para in s.text_frame.paragraphs:
                    for r in para.runs:
                        sz = r.font.size.pt if r.font.size else None
                        col = ""
                        try:
                            if r.font.color and r.font.color.type is not None: col = str(r.font.color.rgb)
                        except Exception: pass
                        fonts.append((r.text[:18], sz, r.font.bold, col))
            txt = s.text_frame.text.replace("\n", " / ").strip() if s.has_text_frame else ""
            shapes.append({"txt": txt, "box": box, "fill": fill, "fonts": tuple(fonts)})
        out.append(shapes)
    return out

def main():
    a, b = load(sys.argv[1]), load(sys.argv[2])
    for i, (sa, sb) in enumerate(zip(a, b), 1):
        msgs = []
        if len(sa) != len(sb):
            print(f"slide {i}: SHAPE COUNT {len(sa)} -> {len(sb)}"); continue
        # text set diff
        ta, tb = {s["txt"] for s in sa if s["txt"]}, {s["txt"] for s in sb if s["txt"]}
        for t in ta - tb: msgs.append(f"  text build-only: {t[:90]}")
        for t in tb - ta: msgs.append(f"  text DELIV-only: {t[:90]}")
        # positional / fill / font diffs (index-aligned)
        for (x, y) in zip(sa, sb):
            if None in x["box"] or None in y["box"]: continue
            d = max(abs(x["box"][k]-y["box"][k]) for k in range(4))
            if d > 4: msgs.append(f"  moved/resized '{(x['txt'] or 'shape')[:22]}' {x['box']} -> {y['box']} (Δ{d}px)")
            if x["fill"] != y["fill"]: msgs.append(f"  fill '{(x['txt'] or 'shape')[:20]}' {x['fill']} -> {y['fill']}")
            for (ra, rb) in zip(x["fonts"], y["fonts"]):
                if (ra[1], ra[2], ra[3]) != (rb[1], rb[2], rb[3]):
                    msgs.append(f"  font '{ra[0]}' size {ra[1]}->{rb[1]} bold {ra[2]}->{rb[2]} col {ra[3]}->{rb[3]}")
        if msgs:
            print(f"slide {i}:"); print("\n".join(msgs))
    print("\nNote: PowerPoint merging pre-wrapped paragraphs into one auto-wrap block is usually cosmetic; a uniform +N px group shift or a font/width/wording change is a real edit to bake into build.py.")

if __name__ == "__main__": main()
