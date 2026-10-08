#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate ANIMATION-TEST.pptx — the engine self-test. Hand this to the user to
run as a PowerPoint slideshow BEFORE animating a full deck, to confirm animlib's
<p:timing> output opens with no repair prompt and PLAYS (loops + click builds).
This exact file was confirmed working in PowerPoint.

Expected in slideshow:
 - On open: BOTH dots pulse continuously (no click needed); boxes hidden.
 - Click 1 -> box 1 fades in; click 2 -> box 2; click 3 -> box 3; click 4 -> dark
   box flies up from below. Dots keep pulsing the whole time.

Usage:  python3 anim_selftest.py [out.pptx]
"""
import sys, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import animlib as AN

E = lambda px: Emu(int(px * 9525))

def build(out="ANIMATION-TEST.pptx"):
    prs = Presentation(); prs.slide_width = E(1920); prs.slide_height = E(1080)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(E(120), E(60), E(1680), E(90))
    tb.text_frame.text = ("Engine self-test — both dots pulse from open; "
                          "click 1-4 reveal the boxes. (No repair prompt expected.)")
    r = tb.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(24); r.font.bold = True
    def box(x, c, t):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(260), E(420), E(170))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(c); sh.line.fill.background()
        sh.text_frame.text = t; rr = sh.text_frame.paragraphs[0].runs[0]
        rr.font.size = Pt(22); rr.font.color.rgb = RGBColor.from_string("FFFFFF")
        return sh
    b1 = box(200, "FF7700", "1 · Fade"); b2 = box(760, "5C6DFF", "2 · Fade")
    b3 = box(1320, "2BB673", "3 · Fade")
    fly = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(200), E(560), E(420), E(170))
    fly.fill.solid(); fly.fill.fore_color.rgb = RGBColor.from_string("050A36"); fly.line.fill.background()
    fly.text_frame.text = "4 · Fly up"; fr = fly.text_frame.paragraphs[0].runs[0]
    fr.font.size = Pt(22); fr.font.color.rgb = RGBColor.from_string("FFFFFF")
    def dot(x, y, d, c):
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(d), E(d))
        o.fill.solid(); o.fill.fore_color.rgb = RGBColor.from_string(c); o.line.fill.background(); return o
    dA = dot(1000, 560, 110, "FF7700"); dB = dot(1450, 575, 80, "E63946")
    AN.animate(s,
        clicks=[[AN.eff(b1, 'fade')], [AN.eff(b2, 'fade')], [AN.eff(b3, 'fade')],
                [AN.eff(fly, 'fly', dir='u')]],
        autos=[AN.eff(dA, 'pulse', dur=700, loop=True),
               AN.eff(dB, 'pulse', dur=500, loop=True)])
    prs.save(out); print("saved", out)

if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "ANIMATION-TEST.pptx")
